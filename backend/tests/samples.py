"""Tiny synthetic files of each accepted type, built in memory."""

import io
import zipfile

from pypdf import PdfWriter


def pdf(pages: int = 1) -> bytes:
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=595, height=842)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def _office(main_part: str, extra: dict[str, str] | None = None) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr(main_part, "<doc/>")
        for name, content in (extra or {}).items():
            archive.writestr(name, content)
    return out.getvalue()


def docx() -> bytes:
    return _office("word/document.xml")


def xlsx() -> bytes:
    return _office("xl/workbook.xml")


def docm() -> bytes:
    return _office("word/document.xml", {"word/vbaProject.bin": "macro"})


PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64
CSV = b"item,ordered,delivered\nWidget,100,80\n"


def text_pdf(pages: list[list[str]]) -> bytes:
    """A real PDF with a text layer: one list of lines per page (Helvetica, A4)."""
    objects: list[bytes] = []
    page_ids = [4 + 2 * i for i in range(len(pages))]
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{pid} 0 R" for pid in page_ids)
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for lines in pages:
        content_id = len(objects) + 2
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>".encode()
        )
        ops = ["BT", "/F1 12 Tf", "72 770 Td", "16 TL"]
        for line in lines:
            escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            ops.append(f"({escaped}) Tj T*")
        ops.append("ET")
        stream = "\n".join(ops).encode()
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref,
    )
    return bytes(out)


def real_docx() -> bytes:
    import docx

    document = docx.Document()
    document.add_paragraph("Purchase Order PO-1001")
    document.add_paragraph("")
    table = document.add_table(rows=2, cols=3)
    for i, text in enumerate(("Item", "Qty ordered", "Unit")):
        table.rows[0].cells[i].text = text
    for i, text in enumerate(("Widget", "100", "units")):
        table.rows[1].cells[i].text = text
    document.add_paragraph("Delivery by 2026-10-30")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def real_xlsx() -> bytes:
    from datetime import date

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Stock count"
    sheet.append(["Item", "Ordered", "Delivered", "Date"])
    sheet.append(["Widget", 100, 80.0, date(2026, 10, 16)])
    sheet.append([None, None, None, None])
    sheet.append(["Bolt", 0.1 + 0.2, 3, None])
    second = workbook.create_sheet("Notes")
    second.append(["Checked by", "Sarah"])
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()
